"""Validated approved workbook → reviewed model → native WinWatt XML/WWP.

This module deliberately refuses to manufacture U values, wall heights or
unreviewed boundary geometry.  A syntactically valid but semantically empty
WWP is not reported as a successful build.
"""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from winwatt_automation.certificates.native_xml import compile_native_xml
from winwatt_automation.certificates.validation import validate_native_readback


def _number(value: Any) -> float | None:
    if value is None or str(value).strip() == "": return None
    try: return float(str(value).replace(" ", "").replace(",", "."))
    except ValueError: return None


def _canonical_rows(workbook: Path) -> list[dict[str, Any]]:
    book=load_workbook(workbook,read_only=True,data_only=True)
    if "Canonical project" not in book.sheetnames: raise ValueError("approved_project.xlsx is missing the Canonical project sheet")
    rows=book["Canonical project"].iter_rows(values_only=True); headers=[str(v) if v is not None else "" for v in next(rows)]
    result=[]
    for row in rows:
        item={headers[i]:row[i] if i<len(row) else None for i in range(len(headers))}
        if item.get("canonical_value") is None: continue
        raw=item.get("source_columns_json") or "{}"
        try: item["source_columns"]=json.loads(raw)
        except (TypeError,json.JSONDecodeError): item["source_columns"]={}
        result.append(item)
    return result


def canonical_to_model(workbook: Path) -> tuple[dict[str, Any], list[str], dict[str, int]]:
    rows=_canonical_rows(workbook); warnings=[]; counts=defaultdict(int)
    records: dict[tuple[str,str],dict[str,Any]]={}
    for row in rows:
        key=(str(row["entity_type"]),str(row["entity_id"])); records.setdefault(key,{**row["source_columns"]})[str(row["field_name"])]=row["canonical_value"]
        counts[str(row["entity_type"])]+=1
    project_name="Approved WebWatt project"; address=""; heated_area=0.0; heated_volume=0.0
    project_fields=[r for (kind,_),r in records.items() if kind=="project_field"]
    for item in project_fields:
        name=str(item.get("Mező") or "").casefold(); value=item.get("PDF-ből talált érték")
        if "terv tárgya" in name and value: project_name=str(value)
        if "település" in name and value: address=(address+" "+str(value)).strip()
    layers=[]; structures=[]
    for (kind,code), item in records.items():
        if kind!="layer": continue
        layer_name=item.get("Réteg / anyag")
        if not layer_name: continue
        thickness=_number(item.get("Vastagság [cm]"))
        if thickness is None: warnings.append(f"layer {code}/{layer_name}: missing or non-scalar thickness")
        layers.append({"structure":code,"sequence":_number(item.get("Réteg sorszám")) or 0,"name":str(layer_name),"thickness_cm":thickness or 0,"density_kgm3":None,"lambda_wmk":None,"heat_capacity_kjkgk":None})
    known_structure=set()
    for layer in layers:
        code=layer["structure"]
        if code in known_structure: continue
        known_structure.add(code)
        structures.append({"name":code,"type":"külső fal","u_effective":None})
    rooms=[]
    for (kind,code), item in records.items():
        if kind!="room": continue
        area=_number(item.get("Alapterület [m²]")); height=_number(item.get("Belmagasság [m]"))
        if area is None: continue
        if height is None: warnings.append(f"room {code}: approved height missing (non-numeric markings such as 1.90 m headroom are not a height)")
        rooms.append({"name":code,"area_m2":area,"height_m":height or 0.0,"volume_m3":area*(height or 0.0),"building":project_name})
        heated_area+=area; heated_volume+=area*(height or 0.0)
    boundaries=[]
    for (kind,code), item in records.items():
        if kind!="boundary": continue
        area=_number(item.get("A [m²]")); x=_number(item.get("x [m]")); y=_number(item.get("y [m]")); structure=str(item.get("Szerkezet / rétegrend") or "")
        if not area or not x or not y: warnings.append(f"boundary {code}: approved x/y/A geometry incomplete")
        if x and y and area and abs(x*y-area)>0.01: warnings.append(f"boundary {code}: x*y differs from A")
        if not structure: warnings.append(f"boundary {code}: structure missing")
        room=str(item.get("Helyiségkód") or code.rsplit("-",1)[0])
        boundary_type=str(item.get("Határoló típusa") or "")
        win_type="külső fal" if "fal" in boundary_type.casefold() else "külső tető" if "tető" in boundary_type.casefold() else "talajon fekvő padló" if "padló" in boundary_type.casefold() else "külső fal"
        boundaries.append({"room":room,"name":code,"winwatt_type":win_type,"area_m2":area or 0.0,"x_m":x,"y_m":y,"u_effective":_number(item.get("U [W/m²K]")),"azimuth_deg":_number(item.get("Tájolás [°]")) or 0,"slope":"függőleges"})
    referenced={str(b["room"]) for b in boundaries}; room_names={str(r["name"]) for r in rooms}
    for name in sorted(referenced-room_names): warnings.append(f"boundary refers to absent approved room {name}")
    model={"project":{"name":project_name,"address":address,"heated_area_m2":heated_area,"heated_volume_m3":heated_volume,"require_wall_xy":True},"buildings":[{"name":project_name,"address":address}],"structures":structures,"layers":layers,"rooms":rooms,"boundaries":boundaries}
    return model,warnings,dict(counts)


def build(workbook: Path, output: Path, *, template_xml: Path | None = None, execute_winwatt: bool = False) -> dict[str, Any]:
    model,warnings,counts=canonical_to_model(workbook); output=output.resolve(); output.parent.mkdir(parents=True,exist_ok=True)
    model_path=output.with_suffix(".canonical.json"); model_path.write_text(json.dumps(model,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    blockers=list(warnings)
    if not counts:
        blockers.append("Canonical project contains no accepted or edited candidates")
    if not model["rooms"]:
        blockers.append("No approved rooms")
    if not model["boundaries"]:
        blockers.append("No approved boundaries")
    blockers.extend([f"structure {item['name']}: U value missing" for item in model["structures"] if item.get("u_effective") is None])
    report={"input":str(workbook.resolve()),"output":str(output),"model":str(model_path),"imported_objects":counts,"generated":{"rooms":len(model["rooms"]),"structures":len(model["structures"]),"layers":len(model["layers"]),"boundaries":len(model["boundaries"])},"mapping_warnings":warnings,"unresolved":blockers,"validation_errors":[],"provenance":"only accepted/edited records from Canonical project were included","wwp_created":False}
    if blockers:
        report["validation_errors"]=["No WWP generated: approved data is semantically incomplete."]
        return report
    if template_xml is None: raise ValueError("template_xml is required to compile approved data")
    xml_path=output.with_suffix(".xml"); report["xml_compile"]=compile_native_xml(model_path,template_xml,xml_path)
    if execute_winwatt:
        from winwatt_automation.services.winwatt_service import WinWattService
        from winwatt_automation.services.xml_native_service import NativeXmlService
        service=WinWattService(); service.create_empty_project(output.with_suffix(".seed.wwp")); NativeXmlService().import_xml(xml_path); service.save_project_as(output); readback=output.with_suffix(".readback.xml"); NativeXmlService().export_xml(readback); report["roundtrip"]=validate_native_readback(model_path,readback); report["wwp_created"]=True
    else:
        report["validation_errors"]=["Native WinWatt execution was not requested; XML was compiled but no WWP was claimed."]
    return report
