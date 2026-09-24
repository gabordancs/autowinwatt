"""Adapters between source workbooks/JSON and the review domain.

No source spreadsheet is mutated.  Candidate IDs are deterministic from the
sheet, source row and reviewed field, so an interrupted review can resume.
"""
from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .domain import DocumentEvidence, ReviewCandidate

SOURCE_TABS = ("Projekt alapadatok", "Rétegrendek", "Helyiségek", "Külső határolók")


def _text(value: object) -> str | None:
    return None if value is None or str(value).strip() == "" else str(value).strip()


def _unit(header: str) -> str | None:
    if "[m²]" in header: return "m²"
    if "[m³]" in header: return "m³"
    if "[m]" in header: return "m"
    if "[cm]" in header: return "cm"
    if "[mm]" in header: return "mm"
    if "[°]" in header: return "°"
    if "[W/m²K]" in header: return "W/m²K"
    return None


def _candidate_id(sheet: str, row: int, field: str) -> str:
    return "rc-" + sha256(f"{sheet}|{row}|{field}".encode("utf-8")).hexdigest()[:20]


def _source_page(source_pdf: str | None) -> int | None:
    # Sheet identifiers (É-02B etc.) are intentionally retained as sheet_no;
    # without a real page number the UI opens the document at page 1.
    return 1 if source_pdf else None


def _row_candidates(sheet: str, row_no: int, values: dict[str, Any]) -> list[ReviewCandidate]:
    source_pdf = _text(values.get("PDF-forrás") or values.get("PDF-forrás / bizonyíték"))
    sheet_no = source_pdf
    method = _text(values.get("Hogyan találtam meg?")) or "spreadsheet_import"
    entity_id = _text(values.get("Helyiségkód") or values.get("Elemkód") or values.get("Kód") or values.get("Mező") or values.get("Csoport"))
    if not entity_id or entity_id in {"MÓDSZER", "DEFINÍCIÓ", "KALIBRÁCIÓ"}:
        return []
    entity_type = {"Projekt alapadatok": "project_field", "Rétegrendek": "layer", "Helyiségek": "room", "Külső határolók": "boundary"}[sheet]
    fields: list[str] = []
    if sheet == "Projekt alapadatok": fields = ["PDF-ből talált érték"]
    elif sheet == "Rétegrendek": fields = ["Réteg / anyag", "Vastagság [cm]"]
    elif sheet == "Helyiségek": fields = ["Helyiség", "Alapterület [m²]", "Belmagasság / jelölés"]
    elif sheet == "Külső határolók": fields = ["Határoló típusa", "Szerkezet / rétegrend", "x [m]", "y [m]", "A [m²]", "Nyílászáró", "Nyílászáró x [m]", "Nyílászáró y [m]"]
    evidence = DocumentEvidence(source_pdf=source_pdf, page=_source_page(source_pdf), sheet_no=sheet_no, extraction_method="spreadsheet_import:" + method[:80], confidence=0.75 if source_pdf else 0.0)
    output = []
    for field in fields:
        value = values.get(field)
        if _text(value) is None:
            continue
        output.append(ReviewCandidate(
            candidate_id=_candidate_id(sheet, row_no, field), entity_type=entity_type, entity_id=entity_id,
            field_name=field, machine_value=value, unit=_unit(field), evidence=evidence,
            source_sheet=sheet, source_row=row_no, source_columns=values,
        ))
    return output


def import_workbook(path: Path) -> list[ReviewCandidate]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    candidates: list[ReviewCandidate] = []
    for sheet in SOURCE_TABS:
        if sheet not in workbook.sheetnames:
            continue
        rows = workbook[sheet].iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        for row_no, row in enumerate(rows, start=2):
            values = {header[index]: row[index] if index < len(row) else None for index in range(len(header)) if header[index]}
            candidates.extend(_row_candidates(sheet, row_no, values))
    return candidates


def import_json(payload: list[dict[str, Any]]) -> list[ReviewCandidate]:
    return [ReviewCandidate.from_dict(item) for item in payload]
