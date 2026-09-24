"""Approved-workbook export. Rejected/unresolved values never become canonical."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from .domain import ReviewStatus
from .persistence import ReviewStore


AUDIT_HEADERS = ["candidate_id", "entity_type", "entity_id", "field_name", "machine_value", "reviewed_value", "canonical_value", "unit", "status", "reviewer", "reviewed_at", "source_pdf", "page", "sheet_no", "evidence_bbox", "extraction_method", "confidence", "note", "source_sheet", "source_row", "source_columns_json"]


def export_approved_workbook(source: Path, target: Path, store: ReviewStore) -> dict:
    workbook = load_workbook(source)
    for name in ("Review audit", "Canonical project"):
        if name in workbook.sheetnames:
            del workbook[name]
    audit = workbook.create_sheet("Review audit")
    canonical = workbook.create_sheet("Canonical project")
    audit.append(AUDIT_HEADERS)
    canonical.append(AUDIT_HEADERS)
    candidates = store.list()
    accepted = 0
    for item in candidates:
        row = [item.candidate_id, item.entity_type, item.entity_id, item.field_name, item.machine_value, item.reviewed_value, item.canonical_value, item.unit, item.status.value, item.reviewer, item.reviewed_at, item.evidence.source_pdf, item.evidence.page, item.evidence.sheet_no, str(item.evidence.evidence_bbox) if item.evidence.evidence_bbox else None, item.evidence.extraction_method, item.evidence.confidence, item.note, item.source_sheet, item.source_row, json.dumps(item.source_columns, ensure_ascii=False, sort_keys=True)]
        audit.append(row)
        if item.status in {ReviewStatus.ACCEPTED, ReviewStatus.EDITED}:
            canonical.append(row)
            accepted += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return {"path": str(target), "candidates": len(candidates), "canonical_candidates": accepted, "excluded": len(candidates) - accepted}
