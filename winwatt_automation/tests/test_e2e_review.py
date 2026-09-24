from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from winwatt_automation.e2e_review.adapters import import_workbook
from winwatt_automation.e2e_review.domain import DocumentEvidence, ReviewCandidate, ReviewStatus, numeric_value
from winwatt_automation.e2e_review.exporter import export_approved_workbook
from winwatt_automation.e2e_review.pdf_evidence import viewport_for_evidence
from winwatt_automation.e2e_review.persistence import ReviewStore


def candidate() -> ReviewCandidate:
    return ReviewCandidate("c-1", "room", "FSZ-01", "Alapterület [m²]", "24,17", unit="m²", evidence=DocumentEvidence("É-02", page=1))


def test_review_state_transition_preserves_machine_value(tmp_path: Path):
    store=ReviewStore(tmp_path/"review.sqlite"); store.seed([candidate()])
    changed=store.transition("c-1",ReviewStatus.ACCEPTED,reviewer="Gábor")
    assert changed.machine_value=="24,17" and changed.reviewed_value=="24,17" and changed.canonical_value=="24,17"
    assert len(store.audit("c-1"))==2


def test_numeric_edit_and_unit_validation(tmp_path: Path):
    store=ReviewStore(tmp_path/"review.sqlite"); store.seed([candidate()])
    changed=store.transition("c-1",ReviewStatus.EDITED,reviewer="Gábor",reviewed_value=numeric_value("24,25","m²"))
    assert changed.machine_value=="24,17" and changed.canonical_value==24.25
    try: numeric_value("nem szám","m²")
    except ValueError: pass
    else: raise AssertionError("non-numeric physical edit was accepted")


def test_bbox_to_pdf_viewport():
    viewport=viewport_for_evidence(DocumentEvidence("É-02",page=2,evidence_bbox=(10,20,40,60),context_margin=10),100,100)
    assert viewport.page==2 and viewport.crop==(0.0,10,50,70) and not viewport.missing_evidence_location
    missing=viewport_for_evidence(DocumentEvidence("É-02"),100,200)
    assert missing.crop==(0.0,0.0,100,200) and missing.missing_evidence_location


def test_workbook_import_export_and_restart_resume(tmp_path: Path):
    source=tmp_path/"source.xlsx"; book=Workbook(); sheet=book.active; sheet.title="Helyiségek"; sheet.append(["Helyiségkód","Helyiség","Alapterület [m²]","PDF-forrás"]); sheet.append(["FSZ-01","NAPPALI","24,17","É-02"]); book.save(source)
    imported=import_workbook(source); assert {item.field_name for item in imported}=={"Helyiség","Alapterület [m²]"}
    db=tmp_path/"review.sqlite"; store=ReviewStore(db); store.seed(imported); area=next(item for item in imported if item.field_name=="Alapterület [m²]"); store.transition(area.candidate_id,ReviewStatus.ACCEPTED,reviewer="Gábor"); store.close()
    resumed=ReviewStore(db); assert resumed.get(area.candidate_id).status is ReviewStatus.ACCEPTED
    output=tmp_path/"approved_project.xlsx"; result=export_approved_workbook(source,output,resumed); resumed.close()
    assert result["canonical_candidates"]==1
    exported=load_workbook(output,data_only=True); assert "Review audit" in exported.sheetnames and "Canonical project" in exported.sheetnames
